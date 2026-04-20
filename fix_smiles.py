#!/usr/bin/env python3
"""
Fixes empty SMILES in pubchem_cache.json and cascade_cache.json.
Uses stored CIDs to batch-fetch SMILES from PubChem (100 CIDs per request).
Then rebuilds the pubchem_data sheet in the Excel workbook.
"""
import json, time, sys, requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

EXCEL_FILE    = "PhenolDB_estruturado.xlsx"
PUBCHEM_CACHE = "pubchem_cache.json"
CASCADE_CACHE = "cascade_cache.json"
BATCH         = 100

with open(PUBCHEM_CACHE, encoding="utf-8") as f: pc = json.load(f)
with open(CASCADE_CACHE, encoding="utf-8") as f: cc = json.load(f)

# ── collect CID → name mapping ───────────────────────────────────────────────
cid_to_names = {}   # cid_str → [name, ...]
for name, data in {**pc, **cc}.items():
    if data.get("status") == "found" and data.get("cid"):
        cid = str(int(data["cid"]))
        cid_to_names.setdefault(cid, []).append((name, "pc" if name in pc else "cc"))

print(f"Found entries with CID: {len(cid_to_names)}")

# ── batch fetch SMILES from PubChem by CID ───────────────────────────────────
cids = list(cid_to_names.keys())
batches = [cids[i:i+BATCH] for i in range(0, len(cids), BATCH)]
smiles_map = {}   # cid_str → {isomeric_smiles, canonical_smiles}

print(f"Fetching SMILES in {len(batches)} batches…")
for i, batch in enumerate(batches, 1):
    url = (
        "https://pubchem.ncbi.nlm.nih.gov/rest/pug/compound/cid/"
        + ",".join(batch)
        + "/property/IsomericSMILES,CanonicalSMILES,IUPACName,MolecularFormula,MolecularWeight/JSON"
    )
    try:
        r = requests.get(url, timeout=30)
        if r.status_code == 200:
            for p in r.json()["PropertyTable"]["Properties"]:
                cid = str(p["CID"])
                smiles_map[cid] = {
                    "isomeric_smiles":   p.get("SMILES", ""),
                    "canonical_smiles":  p.get("ConnectivitySMILES", ""),
                    "iupac_name":        p.get("IUPACName", ""),
                    "molecular_formula": p.get("MolecularFormula", ""),
                    "molecular_weight":  str(p.get("MolecularWeight", "")),
                }
        print(f"  Batch {i}/{len(batches)}: {r.status_code}, {len(smiles_map)} total CIDs resolved")
    except Exception as e:
        print(f"  Batch {i} error: {e}")
    time.sleep(0.3)

# ── update caches ─────────────────────────────────────────────────────────────
updated_pc, updated_cc = 0, 0
for cid, names_list in cid_to_names.items():
    if cid not in smiles_map:
        continue
    for name, source in names_list:
        target = pc if source == "pc" else cc
        if name in target:
            target[name].update(smiles_map[cid])
            if source == "pc": updated_pc += 1
            else: updated_cc += 1

with open(PUBCHEM_CACHE, "w", encoding="utf-8") as f:
    json.dump(pc, f, ensure_ascii=False, indent=2)
with open(CASCADE_CACHE, "w", encoding="utf-8") as f:
    json.dump(cc, f, ensure_ascii=False, indent=2)

print(f"Updated: {updated_pc} PubChem + {updated_cc} cascade entries")

# ── verify ────────────────────────────────────────────────────────────────────
has_smiles = sum(1 for v in {**pc, **cc}.values()
                 if v.get("status")=="found" and v.get("isomeric_smiles"))
print(f"Entries with SMILES now: {has_smiles}")

# ── rebuild pubchem_data sheet ────────────────────────────────────────────────
print("Rebuilding pubchem_data sheet…")

def best_result(name):
    if pc.get(name, {}).get("status") == "found": return pc[name]
    if cc.get(name, {}).get("status") == "found": return cc[name]
    return None

H_FILL = PatternFill("solid", fgColor="1F4E79")
H_FONT = Font(bold=True, color="FFFFFF", size=10)
A_FILL = PatternFill("solid", fgColor="D6E4F0")
BD = Border(
    left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),  bottom=Side(style="thin", color="CCCCCC"),
)

wb = openpyxl.load_workbook(EXCEL_FILE)
if "pubchem_data" in wb.sheetnames: del wb["pubchem_data"]

ws_comp = wb["compounds"]
unique_names = []
seen = set()
for row in ws_comp.iter_rows(min_row=2, values_only=True):
    name = str(row[4]).strip() if row[4] else ""
    if name and name not in seen:
        seen.add(name); unique_names.append(name)

HEADERS = [
    "molecule_name", "status", "source", "cid",
    "isomeric_smiles", "canonical_smiles",
    "iupac_name", "molecular_formula", "molecular_weight",
    "smiles_confidence",
]

ws_out = wb.create_sheet(title="pubchem_data")
for ci, h in enumerate(HEADERS, 1):
    c = ws_out.cell(row=1, column=ci, value=h)
    c.font = H_FONT; c.fill = H_FILL; c.border = BD
    c.alignment = Alignment(horizontal="center", vertical="center")
ws_out.row_dimensions[1].height = 28
ws_out.freeze_panes = "A2"

for ri, name in enumerate(unique_names, 2):
    info = best_result(name) or {}
    status   = info.get("status", "not_found")
    source   = info.get("source", "")
    has_smi  = bool(info.get("isomeric_smiles") or info.get("canonical_smiles"))
    confidence = "confirmed" if (status == "found" and has_smi) else \
                 "found_no_smiles" if status == "found" else "not_found"
    row_data = [
        name, status, source,
        info.get("cid", ""),
        info.get("isomeric_smiles", ""),
        info.get("canonical_smiles", ""),
        info.get("iupac_name", ""),
        info.get("molecular_formula", ""),
        info.get("molecular_weight", ""),
        confidence,
    ]
    fill = A_FILL if ri % 2 == 0 else PatternFill()
    for ci, val in enumerate(row_data, 1):
        c = ws_out.cell(row=ri, column=ci, value=val)
        c.fill = fill; c.border = BD
        c.alignment = Alignment(vertical="top")

for col in ws_out.columns:
    letter = get_column_letter(col[0].column)
    best = max((len(str(c.value or "")) for c in col), default=10)
    ws_out.column_dimensions[letter].width = min(best + 4, 70)

wb.save(EXCEL_FILE)

confirmed = sum(1 for n in unique_names if (best_result(n) or {}).get("isomeric_smiles") or (best_result(n) or {}).get("canonical_smiles"))
print(f"\nSaved '{EXCEL_FILE}'")
print(f"  Compostos com SMILES: {confirmed} / {len(unique_names)} ({100*confirmed//len(unique_names)}%)")
