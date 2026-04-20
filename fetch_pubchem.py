#!/usr/bin/env python3
"""
Queries PubChem PUG REST API for each unique molecule name in PhenolDB_estruturado.xlsx,
retrieves CID, IsomericSMILES, CanonicalSMILES, IUPACName, MolecularFormula, MolecularWeight,
caches results in pubchem_cache.json, then adds a 'pubchem_data' sheet to the workbook.

Rate limit: 5 requests/second (PubChem guideline).
Resume-safe: already-cached names are skipped.
"""

import json
import time
import sys
import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

EXCEL_FILE  = "PhenolDB_estruturado.xlsx"
CACHE_FILE  = "pubchem_cache.json"
SLEEP       = 0.22   # ~4.5 req/s, safely under 5/s limit
BATCH_SAVE  = 50     # save cache every N requests

PUBCHEM_URL = (
    "https://pubchem.ncbi.nlm.nih.gov/rest/pug/compound/name/{name}"
    "/property/IsomericSMILES,CanonicalSMILES,IUPACName,"
    "MolecularFormula,MolecularWeight/JSON"
)

# ── load cache ───────────────────────────────────────────────────────────────
try:
    with open(CACHE_FILE, encoding="utf-8") as f:
        cache = json.load(f)
    print(f"Cache loaded: {len(cache)} entries")
except FileNotFoundError:
    cache = {}
    print("No cache found, starting fresh.")


def save_cache():
    with open(CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)


# ── collect unique molecule names ────────────────────────────────────────────
print("Reading compound names…")
wb = openpyxl.load_workbook(EXCEL_FILE, read_only=True)
ws = wb["compounds"]

unique_names = []
seen = set()
for row in ws.iter_rows(min_row=2, values_only=True):
    name = str(row[4]).strip() if row[4] else ""
    if name and name not in seen:
        seen.add(name)
        unique_names.append(name)
wb.close()
print(f"  {len(unique_names)} unique molecule names")

# ── query PubChem ────────────────────────────────────────────────────────────
to_fetch = [n for n in unique_names if n not in cache]
print(f"  {len(to_fetch)} still need to be fetched\n")

for idx, name in enumerate(to_fetch, 1):
    url = PUBCHEM_URL.format(name=requests.utils.quote(name))
    try:
        r = requests.get(url, timeout=15)
        if r.status_code == 200:
            props = r.json()["PropertyTable"]["Properties"][0]
            cache[name] = {
                "cid":               props.get("CID", ""),
                "isomeric_smiles":   props.get("IsomericSMILES", ""),
                "canonical_smiles":  props.get("CanonicalSMILES", ""),
                "iupac_name":        props.get("IUPACName", ""),
                "molecular_formula": props.get("MolecularFormula", ""),
                "molecular_weight":  props.get("MolecularWeight", ""),
                "status":            "found",
            }
        elif r.status_code == 404:
            cache[name] = {"status": "not_found"}
        else:
            cache[name] = {"status": f"error_{r.status_code}"}
    except Exception as e:
        cache[name] = {"status": f"exception: {e}"}

    status = cache[name].get("status", "?")
    print(f"  [{idx}/{len(to_fetch)}] {name[:55]:<55}  {status}")

    if idx % BATCH_SAVE == 0:
        save_cache()

    time.sleep(SLEEP)

save_cache()
print(f"\nCache saved: {len(cache)} total entries.")

# ── summary ──────────────────────────────────────────────────────────────────
found     = sum(1 for v in cache.values() if v.get("status") == "found")
not_found = sum(1 for v in cache.values() if v.get("status") == "not_found")
errors    = len(cache) - found - not_found
print(f"  Found: {found}  |  Not found: {not_found}  |  Errors: {errors}")

# ── write pubchem_data sheet into the Excel workbook ─────────────────────────
print("\nUpdating workbook…")
wb = openpyxl.load_workbook(EXCEL_FILE)

# Remove old sheet if re-running
if "pubchem_data" in wb.sheetnames:
    del wb["pubchem_data"]

H_FILL = PatternFill("solid", fgColor="1F4E79")
H_FONT = Font(bold=True, color="FFFFFF", size=10)
A_FILL = PatternFill("solid", fgColor="D6E4F0")
BD     = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)

HEADERS = [
    "molecule_name", "cid", "status",
    "isomeric_smiles", "canonical_smiles",
    "iupac_name", "molecular_formula", "molecular_weight",
]

ws_out = wb.create_sheet(title="pubchem_data")

for ci, h in enumerate(HEADERS, 1):
    c = ws_out.cell(row=1, column=ci, value=h)
    c.font = H_FONT; c.fill = H_FILL; c.border = BD
    c.alignment = Alignment(horizontal="center", vertical="center")
ws_out.row_dimensions[1].height = 28
ws_out.freeze_panes = "A2"

for ri, name in enumerate(unique_names, 2):
    info = cache.get(name, {"status": "not_queried"})
    row_data = [
        name,
        info.get("cid", ""),
        info.get("status", ""),
        info.get("isomeric_smiles", ""),
        info.get("canonical_smiles", ""),
        info.get("iupac_name", ""),
        info.get("molecular_formula", ""),
        info.get("molecular_weight", ""),
    ]
    fill = A_FILL if ri % 2 == 0 else PatternFill()
    for ci, val in enumerate(row_data, 1):
        c = ws_out.cell(row=ri, column=ci, value=val)
        c.fill = fill; c.border = BD
        c.alignment = Alignment(vertical="top")

# auto column width
for col in ws_out.columns:
    letter = get_column_letter(col[0].column)
    best = max((len(str(c.value or "")) for c in col), default=10)
    ws_out.column_dimensions[letter].width = min(best + 4, 60)

wb.save(EXCEL_FILE)
print(f"Saved '{EXCEL_FILE}' with pubchem_data sheet ({len(unique_names)} rows).")
