#!/usr/bin/env python3
"""
SMILES lookup via NCI CIR and COCONUT for compounds not found in PubChem/cascade.

Pipeline per compound:
  1. NCI Chemical Identifier Resolver (CIR) by original name
  2. NCI CIR by name variants (stereo removed, quote standardised)
  3. COCONUT natural products database by name
  Mark remaining as 'not_found'.

Results stored in cir_coconut_cache.json and written to PhenolDB_estruturado.xlsx.
"""

import json, re, time, sys, requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

EXCEL_FILE       = "PhenolDB_estruturado.xlsx"
PUBCHEM_CACHE    = "pubchem_cache.json"
CASCADE_CACHE    = "cascade_cache.json"
CIR_COCONUT_CACHE = "cir_coconut_cache.json"
SLEEP            = 0.2
BATCH_SAVE       = 50
TIMEOUT          = 10

CIR_URL     = "https://cactus.nci.nih.gov/chemical/structure/{name}/smiles"
COCONUT_URL = "https://coconut.naturalproducts.net/api/search/compounds?query={name}&page=0&size=1"

STEREO_PREFIXES = re.compile(
    r"^(trans[- ]|cis[- ]|\(?[eEzZ]\)?[- ]|α-|β-|α |β |\([-+]\)-|\(±\)-)", re.I
)


def normalise_name(name):
    n = STEREO_PREFIXES.sub("", name).strip()
    n = re.sub(r"\s{2,}", " ", n)
    return n


def name_variants(original):
    variants = []
    norm = normalise_name(original)
    if norm != original:
        variants.append(norm)
    std = original.replace("\u2019", "'").replace("\u201c", '"').replace("\u201d", '"')
    std = std.replace("\u2032", "'").replace("\u2033", '"')
    if std != original:
        variants.append(std)
    # remove trailing isomer/derivative
    cleaned = re.sub(r"\s+(isomer|derivative|analogue|analog)\s*\d*$", "", original, flags=re.I).strip()
    if cleaned != original:
        variants.append(cleaned)
    # lowercase
    variants.append(original.lower())
    return list(dict.fromkeys(variants))


def cir_lookup(name):
    """Query NCI CIR for SMILES by name."""
    url = CIR_URL.format(name=requests.utils.quote(name))
    try:
        r = requests.get(url, timeout=TIMEOUT)
        if r.status_code == 200:
            smiles = r.text.strip().split("\n")[0].strip()
            if smiles and not smiles.startswith("<"):
                return {
                    "cid": "",
                    "isomeric_smiles": smiles,
                    "canonical_smiles": smiles,
                    "iupac_name": "",
                    "molecular_formula": "",
                    "molecular_weight": "",
                    "status": "found",
                }
    except Exception:
        pass
    return None


def coconut_lookup(name):
    """Query COCONUT natural products database by name."""
    url = COCONUT_URL.format(name=requests.utils.quote(name))
    try:
        r = requests.get(url, timeout=TIMEOUT)
        if r.status_code != 200:
            return None
        data = r.json()
        # COCONUT v2 response: {"data": [...], "meta": {...}}
        items = data.get("data") or []
        if not items:
            return None
        mol = items[0]
        smiles = mol.get("canonical_smiles") or mol.get("smiles") or ""
        if not smiles:
            return None
        return {
            "cid": "",
            "isomeric_smiles": mol.get("smiles", smiles),
            "canonical_smiles": smiles,
            "iupac_name": mol.get("iupac_name", ""),
            "molecular_formula": mol.get("molecular_formula", ""),
            "molecular_weight": str(mol.get("molecular_weight", "")),
            "status": "found",
        }
    except Exception:
        pass
    return None


# ── load caches ──────────────────────────────────────────────────────────────

with open(PUBCHEM_CACHE, encoding="utf-8") as f:
    pc = json.load(f)
with open(CASCADE_CACHE, encoding="utf-8") as f:
    cc = json.load(f)

try:
    with open(CIR_COCONUT_CACHE, encoding="utf-8") as f:
        cir_cc = json.load(f)
    print(f"CIR/COCONUT cache loaded: {len(cir_cc)} entries")
except FileNotFoundError:
    cir_cc = {}
    print("No CIR/COCONUT cache, starting fresh.")


def save_cache():
    with open(CIR_COCONUT_CACHE, "w", encoding="utf-8") as f:
        json.dump(cir_cc, f, ensure_ascii=False, indent=2)


def best_result(name):
    if pc.get(name, {}).get("status") == "found":
        return pc[name]
    if cc.get(name, {}).get("status") == "found":
        return cc[name]
    if cir_cc.get(name, {}).get("status") == "found":
        return cir_cc[name]
    return None


# ── collect names that need lookup ───────────────────────────────────────────

not_found = [n for n in pc if best_result(n) is None]
to_process = [n for n in not_found if n not in cir_cc]
print(f"Not found in PubChem+cascade: {len(not_found)}")
print(f"Still to process:             {len(to_process)}\n")

count = 0
for idx, name in enumerate(to_process, 1):
    result = None
    source = ""

    # Step 1: NCI CIR original name
    result = cir_lookup(name)
    time.sleep(SLEEP)
    if result:
        source = f"cir:{name}"

    # Step 2: NCI CIR with name variants
    if not result:
        for variant in name_variants(name):
            result = cir_lookup(variant)
            time.sleep(SLEEP)
            if result:
                source = f"cir_variant:{variant}"
                break

    # Step 3: COCONUT
    if not result:
        result = coconut_lookup(name)
        time.sleep(SLEEP)
        if result:
            source = f"coconut:{name}"

    # Step 4: COCONUT with normalised name
    if not result:
        norm = normalise_name(name)
        if norm != name:
            result = coconut_lookup(norm)
            time.sleep(SLEEP)
            if result:
                source = f"coconut_norm:{norm}"

    if result:
        result["source"] = source
        cir_cc[name] = result
        status = "found"
    else:
        cir_cc[name] = {"status": "not_found", "source": ""}
        status = "not_found"

    count += 1
    print(f"  [{idx}/{len(to_process)}] {name[:55]:<55}  {status}  ({source})")

    if count % BATCH_SAVE == 0:
        save_cache()

save_cache()

found_new = sum(1 for v in cir_cc.values() if v.get("status") == "found")
print(f"\nCIR/COCONUT: {found_new}/{len(cir_cc)} resolved")

# ── update Excel pubchem_data sheet ─────────────────────────────────────────

print("\nUpdating workbook…")
wb = openpyxl.load_workbook(EXCEL_FILE)

if "pubchem_data" in wb.sheetnames:
    del wb["pubchem_data"]

H_FILL = PatternFill("solid", fgColor="1F4E79")
H_FONT = Font(bold=True, color="FFFFFF", size=10)
A_FILL = PatternFill("solid", fgColor="D6E4F0")
BD = Border(
    left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),  bottom=Side(style="thin", color="CCCCCC"),
)

HEADERS = [
    "molecule_name", "status", "source", "cid",
    "isomeric_smiles", "canonical_smiles",
    "iupac_name", "molecular_formula", "molecular_weight",
    "smiles_confidence",
]

ws_comp = wb["compounds"]
unique_names = []
seen = set()
for row in ws_comp.iter_rows(min_row=2, values_only=True):
    name = str(row[4]).strip() if row[4] else ""
    if name and name not in seen:
        seen.add(name)
        unique_names.append(name)

ws_out = wb.create_sheet(title="pubchem_data")
for ci, h in enumerate(HEADERS, 1):
    c = ws_out.cell(row=1, column=ci, value=h)
    c.font = H_FONT; c.fill = H_FILL; c.border = BD
    c.alignment = Alignment(horizontal="center", vertical="center")
ws_out.row_dimensions[1].height = 28
ws_out.freeze_panes = "A2"

for ri, name in enumerate(unique_names, 2):
    info = best_result(name) or {}
    status = info.get("status", "not_found")
    has_smi = bool(info.get("isomeric_smiles") or info.get("canonical_smiles"))
    confidence = "confirmed" if (status == "found" and has_smi) else \
                 "found_no_smiles" if status == "found" else "not_found"
    row_data = [
        name, status, info.get("source", ""),
        info.get("cid", ""),
        info.get("isomeric_smiles", ""),
        info.get("canonical_smiles", ""),
        info.get("iupac_name", ""),
        info.get("molecular_formula", ""),
        info.get("molecular_weight", ""),
        confidence,
    ]
    fill = A_FILL if ri % 2 == 0 else PatternFill()
    for ci, val in enumerate(row_data, 1):
        c = ws_out.cell(row=ri, column=ci, value=val)
        c.fill = fill; c.border = BD
        c.alignment = Alignment(vertical="top")

for col in ws_out.columns:
    letter = get_column_letter(col[0].column)
    best = max((len(str(c.value or "")) for c in col), default=10)
    ws_out.column_dimensions[letter].width = min(best + 4, 70)

wb.save(EXCEL_FILE)

total = len(unique_names)
confirmed = sum(1 for n in unique_names if (best_result(n) or {}).get("isomeric_smiles") or (best_result(n) or {}).get("canonical_smiles"))
print(f"\nSaved '{EXCEL_FILE}'")
print(f"  Compostos com SMILES: {confirmed} / {total} ({100*confirmed//total}%)")
print(f"  Restantes sem SMILES: {total - confirmed}")
