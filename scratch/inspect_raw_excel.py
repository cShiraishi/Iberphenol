import openpyxl
import os

filename = "Base dados_CP.xlsx"
if os.path.exists(filename):
    print "Inspecting " + filename
    wb = openpyxl.load_workbook(filename, read_only=True, data_only=True)
    for sheet in wb.sheetnames:
        print "Sheet: " + sheet
        ws = wb[sheet]
        # Get the first 5 rows to see headers
        rows = list(ws.iter_rows(max_row=5, values_only=True))
        if rows:
            for i, row in enumerate(rows):
                print "  Row " + str(i+1) + ": " + ", ".join([str(c) for c in row if c is not None][:15])
    wb.close()
else:
    print "File not found: " + filename
